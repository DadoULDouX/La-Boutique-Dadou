from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


HEADERS = ["Date", "Type", "Référence", "Article", "Catégorie", "Plateforme", "Montant", "Marge", "Transaction"]


def export_csv(data):
    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(HEADERS)
    for row in data["rows"]:
        writer.writerow([row["date"], row["kind"].title(), row["code"], row["designation"], row.get("category") or "",
                         row.get("platform") or "", f"{row['amount']:.2f}", "" if row.get("margin") is None else f"{row['margin']:.2f}", row.get("transaction_ref") or ""])
    writer.writerow([])
    s = data["summary"]
    writer.writerow(["SYNTHÈSE", "", "", "Achats", "", "", f"{s['purchases']:.2f}"])
    writer.writerow(["", "", "", "Ventes", "", "", f"{s['revenue']:.2f}"])
    writer.writerow(["", "", "", "Marge", "", "", f"{s['margin']:.2f}"])
    return output.getvalue().encode("utf-8")


def export_xlsx(data, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Livre de compte"
    ws.merge_cells("A1:I1")
    ws["A1"] = "LA BOUTIQUE À DADOU"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="22211F")
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.merge_cells("A2:I2")
    ws["A2"] = title
    ws["A2"].font = Font(size=11, color="6B675F")
    ws.append([])
    ws.append(HEADERS)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B56D3B")
    for row in data["rows"]:
        ws.append([row["date"], row["kind"].title(), row["code"], row["designation"], row.get("category") or "",
                   row.get("platform") or "", row["amount"], row.get("margin"), row.get("transaction_ref") or ""])
    end = ws.max_row
    for row in ws.iter_rows(min_row=5, max_row=end, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = '#,##0.00 [$€-fr-FR]'
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{max(4,end)}"
    widths = [13, 12, 14, 32, 18, 16, 15, 15, 22]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    summary = wb.create_sheet("Synthèse")
    summary["A1"] = "SYNTHÈSE"
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="22211F")
    summary.merge_cells("A1:C1")
    summary.append(["Indicateur", "Montant", "Volume"])
    s = data["summary"]
    summary.append(["Achats", s["purchases"], s["purchase_count"]])
    summary.append(["Ventes", s["revenue"], s["sale_count"]])
    summary.append(["Marge brute", s["margin"], ""])
    for c in summary[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="B56D3B")
    for row in range(3, 6):
        summary.cell(row, 2).number_format = '#,##0.00 [$€-fr-FR]'
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 14
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def export_pdf(data, title):
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=14*mm,
                            leftMargin=14*mm, topMargin=14*mm, bottomMargin=14*mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="DadouTitle", parent=styles["Title"], fontName="Helvetica-Bold",
                              fontSize=19, leading=22, textColor=colors.HexColor("#232420"), alignment=0))
    styles.add(ParagraphStyle(name="DadouRight", parent=styles["BodyText"], alignment=TA_RIGHT,
                              fontSize=8, textColor=colors.HexColor("#777970")))
    story = [Paragraph("LA BOUTIQUE À DADOU", styles["DadouTitle"]),
             Paragraph(title, styles["BodyText"]), Spacer(1, 7*mm)]
    summary = data["summary"]
    summary_table = Table([
        ["TOTAL DES ACHATS", "CHIFFRE D’AFFAIRES", "MARGE BRUTE"],
        [f"{summary['purchases']:.2f} EUR", f"{summary['revenue']:.2f} EUR", f"{summary['margin']:.2f} EUR"],
    ], colWidths=[60*mm]*3)
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#232420")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,-1), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,0), 8), ("FONTSIZE", (0,1), (-1,1), 14),
        ("BACKGROUND", (0,1), (-1,1), colors.HexColor("#F4F1EB")),
        ("BOTTOMPADDING", (0,0), (-1,-1), 8), ("TOPPADDING", (0,0), (-1,-1), 8),
        ("GRID", (0,0), (-1,-1), .5, colors.HexColor("#DDD8CF")),
    ]))
    story += [summary_table, Spacer(1, 8*mm)]
    rows = [["Date", "Type", "Référence", "Article", "Catégorie", "Plateforme", "Montant", "Marge"]]
    for row in data["rows"]:
        rows.append([row["date"], row["kind"].title(), row["code"], row["designation"],
                     row.get("category") or "", row.get("platform") or "—",
                     f"{row['amount']:.2f} EUR", "—" if row.get("margin") is None else f"{row['margin']:.2f} EUR"])
    table = Table(rows, repeatRows=1, colWidths=[20*mm,18*mm,23*mm,53*mm,30*mm,26*mm,25*mm,25*mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#B56D3B")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTNAME", (0,1), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), 7.5),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#DDD8CF")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F8F6F1")]),
        ("ALIGN", (6,1), (-1,-1), "RIGHT"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story.append(table)
    doc.build(story)
    return stream.getvalue()
